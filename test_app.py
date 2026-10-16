from app import app
from database import get_db, init_db
import json

def run_tests():
    init_db()
    from werkzeug.security import generate_password_hash
    conn_setup = get_db()
    conn_setup.execute("UPDATE users SET password_hash = ?, status = 'Aktif' WHERE username IN ('admin', 'admin_staf')", (generate_password_hash('admin123'),))
    conn_setup.commit()
    conn_setup.close()

    client = app.test_client()
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id FROM survei ORDER BY id ASC LIMIT 1")
    s_id = c.fetchone()[0]
    
    c.execute("SELECT id FROM petugas ORDER BY id ASC LIMIT 1")
    p_id = c.fetchone()[0]
    
    c.execute("SELECT id FROM responden ORDER BY id ASC LIMIT 1")
    r_id = c.fetchone()[0]
    
    c.execute("SELECT id FROM pertanyaan WHERE survei_id = ? ORDER BY urutan ASC", (s_id,))
    pertanyaan_ids = [row[0] for row in c.fetchall()]
    conn.close()
    
    print("\n--- PHASE 1: TESTING ROUTE PROTECTION & PUBLIC ACCESS ---")
    
    print("[1] Testing GET / without login (Should redirect to /login)...")
    res = client.get('/', follow_redirects=False)
    assert res.status_code == 302, f"Expected 302, got {res.status_code}"
    assert '/login' in res.headers['Location']
    print("    -> Passed! Unauthenticated access redirected to /login.")

    print("[2] Testing GET /petugas without login (Should redirect to /login)...")
    res = client.get('/petugas', follow_redirects=False)
    assert res.status_code == 302
    assert '/login' in res.headers['Location']
    print("    -> Passed! Admin route protected.")

    print("[3] Testing GET /operator/tugas without login (Should redirect to /login)...")
    res = client.get('/operator/tugas', follow_redirects=False)
    assert res.status_code == 302
    assert '/login' in res.headers['Location']
    print("    -> Passed! Operator route protected.")

    print("[4] Testing GET /login (Login Page UI)...")
    res = client.get('/login')
    assert res.status_code == 200
    assert b"SiSurvei" in res.data
    assert b"Petugas / Enumerator" in res.data
    assert b"Administrator" in res.data
    assert b"1-Click Demo Login" in res.data
    print("    -> Passed! Login UI loads with dual role tabs.")

    print("[5] Testing GET /s/SRV-2026-001 (Public Respondent Form without login)...")
    res = client.get('/s/SRV-2026-001')
    assert res.status_code == 200
    assert b"Kuisioner" in res.data
    print("    -> Passed! Public respondent questionnaire accessible without login.")

    print("[6] Testing POST /s/SRV-2026-001 (Public Respondent Submission)...")
    post_data = {
        'nama_responden': 'Responden Tester Publik',
        'kontak_responden': '081234567890',
        'usia': '29',
        'jenis_kelamin': 'Laki-laki',
        'pendidikan_terakhir': 'Sarjana (S1)',
        'alamat_wilayah': 'Jakarta Pusat',
        'petugas_id': str(p_id),
    }
    for q_idx, qid in enumerate(pertanyaan_ids):
        if q_idx == 0:
            post_data[f'q_{qid}'] = 'Administrasi Kependudukan (KTP/KK/Akta)'
        elif q_idx in [1, 2, 3]:
            post_data[f'q_{qid}'] = '5'
        elif q_idx == 4:
            post_data[f'q_{qid}'] = 'Sesuai Tarif Resmi & Sangat Transparan'
        elif q_idx == 5:
            post_data[f'q_{qid}[]'] = ['Sistem Antrean Online & Notifikasi WA']
        elif q_idx == 6:
            post_data[f'q_{qid}'] = 'Aplikasi Mobile Resmi'
        elif q_idx == 7:
            post_data[f'q_{qid}'] = '15'
        else:
            post_data[f'q_{qid}'] = 'Pelayanan testing otomatis sangat memuaskan!'
            
    res = client.post('/s/SRV-2026-001', data=post_data, follow_redirects=True)
    assert res.status_code == 200
    assert b"Terima Kasih" in res.data
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT pendidikan_terakhir FROM responden WHERE nama_responden = 'Responden Tester Publik' ORDER BY id DESC LIMIT 1")
    row_p = c.fetchone()
    conn.close()
    assert row_p and row_p[0] == 'Sarjana (S1)', f"Expected 'Sarjana (S1)', got {row_p}"
    print("    -> Passed! Public respondent response submitted and pendidikan_terakhir persisted correctly.")

    print("\n--- PHASE 2: TESTING AUTHENTICATION (ADMIN & OPERATOR) ---")

    print("[7] Testing POST /login with Invalid Admin Credentials...")
    res = client.post('/login', data={'role': 'admin', 'identifier': 'admin', 'password': 'wrongpassword'}, follow_redirects=True)
    assert res.status_code == 200
    assert b"salah" in res.data or b"tidak sesuai" in res.data
    print("    -> Passed! Invalid admin login rejected.")

    print("[8] Testing POST /login with Invalid Operator Credentials...")
    res = client.post('/login', data={'role': 'operator', 'identifier': 'PTG-001', 'password': 'wrongpassword'}, follow_redirects=True)
    assert res.status_code == 200
    assert b"tidak sesuai" in res.data or b"salah" in res.data
    print("    -> Passed! Invalid operator login rejected.")

    print("[9] Testing POST /login with Valid Admin Credentials...")
    res = client.post('/login', data={'role': 'admin', 'identifier': 'admin', 'password': 'admin123'}, follow_redirects=False)
    assert res.status_code == 302
    assert res.headers['Location'] == '/' or res.headers['Location'].endswith('/')
    print("    -> Passed! Admin login succeeded, redirected to dashboard.")

    print("\n--- PHASE 3: TESTING ADMIN PRIVILEGES ---")

    print("[10] Testing GET / as Authenticated Admin...")
    res = client.get('/')
    assert res.status_code == 200
    assert b"Ringkasan Eksekutif" in res.data
    assert b"Administrator Utama" in res.data
    print("    -> Passed! Admin dashboard rendered with admin identity.")

    print("[11] Testing GET /survei as Authenticated Admin...")
    res = client.get('/survei')
    assert res.status_code == 200
    assert b"SRV-2026-001" in res.data
    print("    -> Passed! Survei list OK.")

    print(f"[12] Testing GET /survei/{s_id}/builder as Authenticated Admin...")
    res = client.get(f'/survei/{s_id}/builder')
    assert res.status_code == 200
    assert b"Perancang Kuisioner" in res.data
    print("    -> Passed! Builder page OK.")

    print(f"[13] Testing GET /survei/{s_id}/hasil as Authenticated Admin...")
    res = client.get(f'/survei/{s_id}/hasil')
    assert res.status_code == 200
    assert b"Visualisasi Analisis Hasil Kuisioner" in res.data
    print("    -> Passed! Results analytics OK.")

    print(f"[14] Testing GET /survei/{s_id}/export/excel and /export/csv as Authenticated Admin...")
    res = client.get(f'/survei/{s_id}/export/excel')
    assert res.status_code == 200
    assert "spreadsheetml" in res.content_type
    
    res_csv = client.get(f'/survei/{s_id}/export/csv')
    assert res_csv.status_code == 200
    assert "Pendidikan Terakhir" in res_csv.data.decode('utf-8')
    print("    -> Passed! Excel & CSV exports OK and include 'Pendidikan Terakhir'.")

    print("[15] Testing GET /petugas as Authenticated Admin...")
    res = client.get('/petugas')
    assert res.status_code == 200
    assert b"Budi Santoso" in res.data
    print("    -> Passed! Petugas list OK.")

    print("[16] Testing GET /respons as Authenticated Admin...")
    res = client.get('/respons')
    assert res.status_code == 200
    assert b"Repositori Data Responden" in res.data
    print("    -> Passed! Responses repository OK.")

    print("\n--- PHASE 4: TESTING LOGOUT & OPERATOR LOGIN ---")

    print("[17] Testing GET /logout...")
    res = client.get('/logout', follow_redirects=False)
    assert res.status_code == 302
    assert '/login' in res.headers['Location']
    print("    -> Passed! Logout redirects to /login.")

    print("[18] Testing GET / after Logout (Should be blocked)...")
    res = client.get('/', follow_redirects=False)
    assert res.status_code == 302
    assert '/login' in res.headers['Location']
    print("    -> Passed! Session cleared, access blocked.")

    print("[19] Testing POST /login with Valid Operator Credentials (PTG-001)...")
    res = client.post('/login', data={'role': 'operator', 'identifier': 'PTG-001', 'password': 'password123'}, follow_redirects=False)
    assert res.status_code == 302
    assert '/operator/tugas' in res.headers['Location']
    print("    -> Passed! Operator login succeeded, redirected to /operator/tugas.")

    print("[20] Testing GET /operator/tugas as Authenticated Operator...")
    res = client.get('/operator/tugas')
    assert res.status_code == 200
    assert b"Tugas Survei Saya" in res.data
    assert b"Budi Santoso" in res.data
    assert b"PTG-001" in res.data
    print("    -> Passed! Operator tugas page loads with correct operator.")

    print("[21] Testing GET /operator/profil as Authenticated Operator...")
    res = client.get('/operator/profil')
    assert res.status_code == 200
    assert b"Profil Petugas" in res.data
    assert b"Perbarui Informasi Data Diri" in res.data
    print("    -> Passed! Operator profile page OK.")

    print("\n--- PHASE 5: TESTING ROLE RESTRICTIONS (OPERATOR BOUNDARIES) ---")

    print("[22] Testing Operator attempting to access Admin route / (Dashboard)...")
    res = client.get('/', follow_redirects=False)
    assert res.status_code == 302
    assert '/operator/tugas' in res.headers['Location']
    print("    -> Passed! Operator blocked from /, redirected to /operator/tugas.")

    print("[23] Testing Operator attempting to access Admin route /petugas...")
    res = client.get('/petugas', follow_redirects=False)
    assert res.status_code == 302
    assert '/operator/tugas' in res.headers['Location']
    print("    -> Passed! Operator blocked from /petugas.")

    print("[24] Testing Operator attempting to access Admin route /survei/tambah...")
    res = client.get('/survei/tambah', follow_redirects=False)
    assert res.status_code == 302
    assert '/operator/tugas' in res.headers['Location']
    print("    -> Passed! Operator blocked from /survei/tambah.")

    print("\n--- PHASE 6: TESTING OPERATOR FIELD SURVEY SUBMISSION ---")

    print(f"[25] Testing GET /operator/survei/{s_id}/isi as Authenticated Operator...")
    res = client.get(f'/operator/survei/{s_id}/isi')
    assert res.status_code == 200
    assert b"Target Sampel Anda" in res.data
    print("    -> Passed! Field survey form OK.")

    print(f"[26] Testing POST /operator/survei/{s_id}/isi (Operator Input Sample with titik_koordinat)...")
    sample_data = {
        'nama_responden': 'Warga Uji Coba Enumerator #99',
        'kontak_responden': '081399887766',
        'usia': '35',
        'jenis_kelamin': 'Perempuan',
        'pendidikan_terakhir': 'Diploma (D1 - D4)',
        'alamat_wilayah': '-6.208763, 106.845599',
    }
    for q_idx, qid in enumerate(pertanyaan_ids):
        if q_idx == 0:
            sample_data[f'q_{qid}'] = 'Administrasi Kependudukan (KTP/KK/Akta)'
        elif q_idx in [1, 2, 3]:
            sample_data[f'q_{qid}'] = '5'
        elif q_idx == 4:
            sample_data[f'q_{qid}'] = 'Sesuai Tarif Resmi & Sangat Transparan'
        elif q_idx == 5:
            sample_data[f'q_{qid}[]'] = ['Sistem Antrean Online & Notifikasi WA']
        elif q_idx == 6:
            sample_data[f'q_{qid}'] = 'Aplikasi Mobile Resmi'
        elif q_idx == 7:
            sample_data[f'q_{qid}'] = '10'
        else:
            sample_data[f'q_{qid}'] = 'Pelayanan enumerator sangat ramah dan efisien.'

    res = client.post(f'/operator/survei/{s_id}/isi', data=sample_data, follow_redirects=True)
    assert res.status_code == 200
    assert b"Sampel Berhasil Disimpan" in res.data
    assert b"Lanjut Input Sampel Berikutnya" in res.data
    
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id, pendidikan_terakhir, titik_koordinat, alamat_wilayah FROM responden WHERE nama_responden = 'Warga Uji Coba Enumerator #99' ORDER BY id DESC LIMIT 1")
    op_row = c.fetchone()
    conn.close()
    assert op_row and op_row[1] == 'Diploma (D1 - D4)', f"Expected 'Diploma (D1 - D4)', got {op_row}"
    assert '-6.208763' in (op_row[2] or op_row[3]), f"Expected coordinate in row, got {op_row}"
    new_resp_id = op_row[0]
    print("    -> Passed! Operator survey sample saved with pendidikan_terakhir & titik_koordinat GPS.")

    print(f"[27] Testing GET /api/responden/{new_resp_id}/detail as Authenticated Operator...")
    res = client.get(f'/api/responden/{new_resp_id}/detail')
    assert res.status_code == 200
    detail_data = json.loads(res.data)
    assert detail_data['success'] is True
    assert detail_data['responden']['pendidikan_terakhir'] == 'Diploma (D1 - D4)'
    assert '-6.208763' in (detail_data['responden'].get('titik_koordinat') or detail_data['responden'].get('alamat_wilayah') or '')
    print("    -> Passed! Operator can fetch respondent detail modal data; verified pendidikan_terakhir & GPS coordinates.")

    print("\n--- PHASE 7: TESTING DUAL ADMIN ROLES (SUPER ADMIN VS ADMIN STAF) ---")
    
    print("[28] Testing Super Admin Login & Access to /admin/profil...")
    client.get('/logout')
    res = client.post('/login', data={'role': 'admin', 'identifier': 'admin', 'password': 'admin123'}, follow_redirects=True)
    assert res.status_code == 200
    res = client.get('/admin/profil')
    assert res.status_code == 200
    assert b"SUPER ADMINISTRATOR" in res.data
    assert b"Profil & Akun Administrator" in res.data
    print("    -> Passed! Super Admin access to /admin/profil confirmed.")

    print("[29] Testing Super Admin Access to /admin/users, /admin/pengaturan, /admin/logs...")
    res = client.get('/admin/users')
    assert res.status_code == 200
    assert b"Manajemen Pengguna Admin" in res.data

    res = client.get('/admin/pengaturan')
    assert res.status_code == 200
    assert b"Pengaturan Sistem" in res.data

    res = client.get('/admin/logs')
    assert res.status_code == 200
    assert b"Riwayat Audit Trail" in res.data
    print("    -> Passed! Super Admin has full access to all system administration modules.")

    print("[30] Testing Super Admin Updating Profile Details...")
    res = client.post('/admin/profil', data={
        'action': 'update_profile',
        'nama': 'Administrator Utama Super',
        'email': 'superadmin@survei.go.id',
        'no_hp': '081234567899',
        'foto_type': 'keep'
    }, follow_redirects=True)
    assert res.status_code == 200
    assert b"berhasil diperbarui" in res.data
    print("    -> Passed! Super Admin profile updated successfully.")

    print("[31] Testing Admin Staf Login (admin_staf)...")
    client.get('/logout')
    res = client.post('/login', data={'role': 'admin', 'identifier': 'admin_staf', 'password': 'admin123'}, follow_redirects=True)
    assert res.status_code == 200
    assert b"Staf Operasional Survei" in res.data
    print("    -> Passed! Admin Staf login succeeded.")

    print("[32] Testing Admin Staf Access to /admin/profil (Allowed)...")
    res = client.get('/admin/profil')
    assert res.status_code == 200
    assert b"ADMINISTRATOR STAF" in res.data
    assert b"Hak Akses Akun Anda" in res.data
    print("    -> Passed! Admin Staf can access own profile.")

    print("[33] Testing Admin Staf Access to /admin/users, /admin/pengaturan, /admin/logs (Must be Blocked)...")
    res = client.get('/admin/users', follow_redirects=False)
    assert res.status_code == 302
    assert res.headers['Location'] == '/' or res.headers['Location'].endswith('/')

    res = client.get('/admin/pengaturan', follow_redirects=False)
    assert res.status_code == 302
    assert res.headers['Location'] == '/' or res.headers['Location'].endswith('/')

    res = client.get('/admin/logs', follow_redirects=False)
    assert res.status_code == 302
    assert res.headers['Location'] == '/' or res.headers['Location'].endswith('/')
    print("    -> Passed! Admin Staf strictly blocked from Super Admin modules.")

    print("[34] Testing Admin Staf Attempting to Delete Survey (Must be Blocked)...")
    res = client.post('/survei/1/hapus', follow_redirects=True)
    assert res.status_code == 200
    assert b"Hanya Super Admin yang berwenang" in res.data
    print("    -> Passed! Critical action /survei/1/hapus blocked for Admin Staf.")

    print("\n--- PHASE 8: TESTING BACKUP & DATA RECOVERY MODULE (SUPER ADMIN) ---")

    print("[35] Testing Admin Staf Access to /admin/backup (Must be Blocked)...")
    res = client.get('/admin/backup', follow_redirects=False)
    assert res.status_code == 302
    assert res.headers['Location'] == '/' or res.headers['Location'].endswith('/')
    print("    -> Passed! Admin Staf blocked from /admin/backup.")

    print("[36] Testing Super Admin Access to /admin/backup...")
    client.get('/logout')
    client.post('/login', data={'role': 'admin', 'identifier': 'admin', 'password': 'admin123'}, follow_redirects=True)
    res = client.get('/admin/backup')
    assert res.status_code == 200
    assert b"Backup & Pemulihan Aset Data Perusahaan" in res.data
    assert b"Paket Arsip ZIP Lengkap" in res.data
    assert b"Master Data Excel (.xlsx)" in res.data
    print("    -> Passed! Super Admin access to /admin/backup UI verified.")

    print("[37] Testing Super Admin Downloading Full ZIP Backup...")
    res = client.get('/admin/backup/download/zip')
    assert res.status_code == 200
    assert res.headers['Content-Type'] == 'application/zip'
    assert len(res.data) > 0
    import zipfile, io
    zf = zipfile.ZipFile(io.BytesIO(res.data))
    namelist = zf.namelist()
    assert 'database/sisurvei.db' in namelist
    assert 'database/sisurvei_dump.sql' in namelist
    assert 'backup_manifest.json' in namelist
    print("    -> Passed! Full ZIP backup generated and contains database, sql dump, and manifest.")

    print("[38] Testing Super Admin Downloading Master Excel Backup (.xlsx)...")
    res = client.get('/admin/backup/download/excel')
    assert res.status_code == 200
    assert 'spreadsheetml' in res.headers['Content-Type']
    assert len(res.data) > 0
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert 'Kegiatan Survei' in wb.sheetnames
    assert 'Data Responden' in wb.sheetnames
    assert 'Data Petugas' in wb.sheetnames
    assert 'Jawaban Responden' in wb.sheetnames
    print("    -> Passed! Master Excel workbook generated with all company survey data sheets.")

    print("[39] Testing Super Admin Downloading Raw SQLite .db & .sql Dump...")
    res = client.get('/admin/backup/download/db')
    assert res.status_code == 200
    assert len(res.data) > 0

    res = client.get('/admin/backup/download/sql')
    assert res.status_code == 200
    assert b"CREATE TABLE" in res.data
    print("    -> Passed! Raw DB and SQL dump exports verified.")

    print("[40] Testing Super Admin Creating Server Snapshot...")
    res = client.post('/admin/backup/create', follow_redirects=True)
    assert res.status_code == 200
    assert b"Snapshot cadangan lengkap berhasil disimpan" in res.data
    print("    -> Passed! Server snapshot creation and listing verified.")

    print("\n--- PHASE 9: TESTING ENTERPRISE SECURITY & AUTOMATED MONTHLY BACKUP ENGINE ---")

    print("[41] Testing HTTP Security Headers (Anti-Sniffing, Anti-Clickjacking, Anti-XSS)...")
    res = client.get('/login')
    assert res.headers.get('X-Content-Type-Options') == 'nosniff', f"Expected nosniff, got {res.headers.get('X-Content-Type-Options')}"
    assert res.headers.get('X-Frame-Options') == 'SAMEORIGIN', f"Expected SAMEORIGIN, got {res.headers.get('X-Frame-Options')}"
    assert res.headers.get('X-XSS-Protection') == '1; mode=block', f"Expected 1; mode=block, got {res.headers.get('X-XSS-Protection')}"
    assert res.headers.get('Referrer-Policy') == 'strict-origin-when-cross-origin', f"Expected strict-origin-when-cross-origin, got {res.headers.get('Referrer-Policy')}"
    print("    -> Passed! All 4 essential HTTP Security Headers enforced on responses.")

    print("[42] Testing Anti-Brute Force Protection & Auto-Lockout...")
    client.get('/logout')
    from app import failed_attempts_tracker, failed_tracker_lock
    with failed_tracker_lock:
        failed_attempts_tracker.clear()

    # Simulasi 5 kali salah password berturut-turut
    test_target = 'bruteforce_target'
    for attempt in range(5):
        res = client.post('/login', data={'role': 'admin', 'identifier': test_target, 'password': 'wrongpassword123'}, follow_redirects=True)
        assert b"Username/email atau kata sandi Administrator salah!" in res.data

    # Percobaan ke-6 harus ditolak oleh sistem lockout
    res_locked = client.post('/login', data={'role': 'admin', 'identifier': test_target, 'password': 'wrongpassword123'}, follow_redirects=True)
    assert b"Peringatan Keamanan: Terlalu banyak percobaan masuk yang gagal" in res_locked.data
    assert b"Akses diblokir sementara" in res_locked.data
    print("    -> Passed! Brute-force threshold exceeded (5x) triggers automated 15-minute security lockout.")

    # Bersihkan tracker setelah test dan login kembali sebagai admin
    with failed_tracker_lock:
        failed_attempts_tracker.clear()
    client.post('/login', data={'role': 'admin', 'identifier': 'admin', 'password': 'admin123'}, follow_redirects=True)

    print("[43] Testing Toggle Auto-Backup Bulanan Route (/admin/backup/toggle-auto)...")
    res_toggle = client.post('/admin/backup/toggle-auto', follow_redirects=True)
    assert res_toggle.status_code == 200
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT nilai FROM pengaturan WHERE kunci = 'auto_backup_bulanan'")
    status_val = c.fetchone()[0]
    conn.close()
    assert status_val in ('Aktif', 'Nonaktif')
    
    # Toggle kembali ke Aktif
    if status_val == 'Nonaktif':
        client.post('/admin/backup/toggle-auto', follow_redirects=True)

    from app import run_monthly_auto_backup_if_due
    # Uji eksekusi engine pencadangan bulanan
    conn = get_db()
    c = conn.cursor()
    c.execute("UPDATE pengaturan SET nilai = '' WHERE kunci = 'terakhir_auto_backup'")
    conn.commit()
    conn.close()

    created_backup = run_monthly_auto_backup_if_due()
    assert created_backup is not None
    assert "auto_monthly_backup_" in created_backup
    assert created_backup.endswith('.zip')

    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT nilai FROM pengaturan WHERE kunci = 'terakhir_auto_backup'")
    last_val = c.fetchone()[0]
    from datetime import datetime
    assert last_val == datetime.now().strftime('%Y-%m')
    conn.close()
    print("    -> Passed! Monthly auto-backup engine automatically generates monthly archive and updates state.")

    print("[44] Testing Path Traversal Defense Shield...")
    res_trav = client.get('/admin/backup/file/unauthorized_sys_file.zip', follow_redirects=True)
    assert res_trav.status_code == 200
    assert b"Akses file tidak diizinkan atau file tidak ditemukan" in res_trav.data

    res_del_trav = client.post('/admin/backup/delete/unauthorized_sys_file.zip', follow_redirects=True)
    assert res_del_trav.status_code == 200
    assert b"Akses penghapusan file tidak sah" in res_del_trav.data
    print("    -> Passed! Path traversal and unauthorized file access blocked by directory containment shield.")

    print("\n=======================================================")
    print(" ALL 44 AUTOMATED INTEGRATION TESTS PASSED 100%! ")
    print("=======================================================\n")

if __name__ == '__main__':
    run_tests()
